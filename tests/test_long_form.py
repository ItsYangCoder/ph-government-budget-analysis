"""Conversion and publication regressions for financial correctness and repeatability."""

from contextlib import closing
from decimal import Decimal
import json
from pathlib import Path
import sqlite3

import openpyxl
import pytest

from pondolens.converters.common import Controls, number, record
from pondolens.converters.excel import agency, grdp, population, sector
from pondolens.converters.regional import CAPTIONS, PAGE1, PAGE2, regional
from pondolens.ingestion import digest
from pondolens import long_form


def cell(v):
    return dict(original_value=v, excel_type="n" if isinstance(v, (int, float)) else "s")


def row(sheet, index, values):
    return sheet, index, {col: cell(v) for col, v in values.items()}


def test_numbers_preserve_zero_sign_precision_and_missing_states():
    assert number(0) == (Decimal(0), "ok")
    assert number("-1,234.56789") == (Decimal("-1234.56789"), "ok")
    assert number("1,23")[1] == "invalid_numeric"
    assert number("#DIV/0!")[1] == "source_excel_error"
    assert number("-") == (None, "source_dash")
    assert number(None) == (None, "missing")
    r = record("test", "sheet", 1, 1, cell("123.456"), unit="thousand_PHP", key=[1])
    assert Decimal(r["value_php"]) == Decimal("123456")


@pytest.fixture
def gaa_file(tmp_path, monkeypatch):
    root = tmp_path / "raw"
    source = root / "02_approved_gaa" / "GAA-2024.xlsx"
    source.parent.mkdir(parents=True)
    workbook = openpyxl.Workbook()
    workbook.active.append(["YEAR", "DEPARTMENT", "AGENCY", "PREXC_FPAP_ID", "FUNDCD", "UACS_SOBJ_CD", "AMT"])
    workbook.active.append(["GAA-2024", "01", "001", "100", "01101101", "501", 0])
    workbook.active.append([None, None, None, None, None, None, 0])
    workbook.save(source)
    # Isolate the real GAA converter and loader; production discovery still requires all nine inputs.
    monkeypatch.setattr(long_form, "raw_inputs", lambda raw: [dict(source=source.relative_to(raw).as_posix(),
                                                                 sha256=digest(source), path=str(source))])
    return root, source


def count(output):
    with closing(sqlite3.connect(output / "long_form.sqlite")) as db:
        return db.execute("SELECT count(*), sum(CAST(value_php AS NUMERIC)) FROM validated_approved_budget").fetchone()


def test_rerun_and_revised_source_replace_current_partition(gaa_file, tmp_path):
    raw, source = gaa_file
    output = tmp_path / "out"
    config = dict(gaa_unit="thousand_PHP", gaa_unit_evidence="Synthetic fixture specifies thousand pesos")
    before = digest(source)
    first = long_form.run(raw, output, **config)
    assert first["ready_for_source_tables"]
    assert count(output) == (1, 0)
    second = long_form.run(raw, output, **config)
    assert second["reused_conversion"]
    assert count(output) == (1, 0)
    assert digest(source) == before
    workbook = openpyxl.load_workbook(source)
    workbook.active["G2"] = 123.456
    workbook.active["G3"] = 123.456
    workbook.save(source)
    third = long_form.run(raw, output, **config)
    assert third["run_id"] != first["run_id"]
    assert Path(first["output_directory"]).exists()
    assert count(output) == (1, 123456)


def test_duplicate_revision_is_retained_for_review_and_removes_stale_partition(gaa_file, tmp_path):
    raw, source = gaa_file
    output = tmp_path / "out"
    config = dict(gaa_unit="PHP", gaa_unit_evidence="Synthetic fixture")
    long_form.run(raw, output, **config)
    workbook = openpyxl.load_workbook(source)
    workbook.active.append(["GAA-2024", "01", "001", "100", "01101101", "501", 55])
    workbook.save(source)
    result = long_form.run(raw, output, **config)
    assert result["sources"][0]["duplicate_count"] == 1
    assert result["sources"][0]["counts"]["approved_budget"] == 2
    assert not result["ready_for_source_tables"]
    assert count(output) == (0, None)


def test_unverified_units_block_load_and_corrupt_outputs_are_rebuilt(gaa_file, tmp_path):
    raw, _ = gaa_file
    output = tmp_path / "out"
    blocked = long_form.run(raw, output)
    assert not blocked["ready_for_source_tables"]
    assert count(output) == (0, None)
    config = dict(gaa_unit="PHP", gaa_unit_evidence="Synthetic fixture")
    passed = long_form.run(raw, output, **config)
    file = next(Path(passed["output_directory"]).glob("*.csv"))
    correct = file.read_bytes()
    file.write_text("corrupt")
    rebuilt = long_form.run(raw, output, **config)
    assert not rebuilt["reused_conversion"]
    assert file.read_bytes() == correct
    assert count(output) == (1, 0)
    with closing(sqlite3.connect(output / "long_form.sqlite")) as db:
        db.execute("DELETE FROM approved_budget"); db.commit()
    assert long_form.run(raw, output, **config)["reused_conversion"]
    assert count(output) == (1, 0)
    (Path(rebuilt["output_directory"]) / "manifest.json").write_text("corrupt")
    assert not long_form.run(raw, output, **config)["reused_conversion"]


def pdf_word(text, x0, top, x1, upright=True):
    return dict(text=text, x0=x0, x1=x1, top=top, bottom=top + 5, upright=upright)


def synthetic_pages():
    pages = []
    for page_number, regions in enumerate([PAGE1, PAGE2], 1):
        words = [pdf_word("CY 2024 Thousand Pesos", 0, 1, 80)]
        for index, region in enumerate(regions):
            anchor = 120 + index * 50
            words.append(pdf_word(CAPTIONS[region], anchor - 20, 10, anchor))
        for label, top in [("DEPARTMENTS", 20), ("Agency A", 30), ("GRAND TOTAL", 40)]:
            words.append(pdf_word(label, 0, top, 70))
            for index, region in enumerate(regions):
                anchor = 120 + index * 50
                # A split four-digit amount must be joined before column assignment.
                if region == "TOTAL":
                    amount = "19,019"
                    words.append(pdf_word(amount, anchor - 20, top, anchor))
                else:
                    words.extend([pdf_word("1", anchor - 26, top, anchor - 21),
                                  pdf_word(",001", anchor - 20, top, anchor)])
        words.append(pdf_word("unrelated", 100, 20, 120, upright=False))
        pages.append(dict(page=page_number, words=words))
    return pages


def test_pdf_split_digits_alignment_and_control_totals():
    controls = Controls()
    records = list(regional(synthetic_pages(), controls, {"source": "2024_B6a.pdf"}))
    assert len(records) == 60
    assert all(r["value_status"] == "ok" for r in records)
    assert next(r["value"] for r in records if r["entity"] == "Agency A" and r["region"] == "NCR") == "1001"
    assert not controls.blockers
    assert all(c["passed"] for c in controls.checks)
    pages = synthetic_pages()
    next(w for w in pages[1]["words"] if w["text"] == "Agency A")["text"] = "Wrong Agency"
    with pytest.raises(ValueError, match="align"):
        list(regional(pages, Controls(), {"source": "2024_B6a.pdf"}))


def test_pdf_amount_change_is_detected_by_reconciliation():
    pages = synthetic_pages()
    target = next(w for w in pages[0]["words"] if w["top"] == 30 and w["text"] == "19,019")
    target["text"] = "20,019"
    controls = Controls()
    list(regional(pages, controls, {"source": "2024_B6a.pdf"}))
    assert any(not c["passed"] for c in controls.checks)


def test_sector_actual_program_proposed_and_units_remain_separate():
    rows = []
    sectors = ["ECONOMIC SERVICES", "SOCIAL SERVICES", "DEFENSE", "GENERAL PUBLIC SERVICES", "NET LENDING", "DEBT SERVICE"]
    suffixes = ["", ", PERCENT DISTRIBUTION", ", GROWTH RATES", ", AS A PERCENTAGE OF GDP",
                ", PER CAPITA", ", REAL LEVELS", ", REAL LEVELS GROWTH RATES"]
    for section, suffix in enumerate(suffixes):
        base = 1 + section * 20
        rows.extend([row("S", base, {1: "NATIONAL GOVERNMENT EXPENDITURES, BY SECTOR" + suffix}),
                     row("S", base + 2, {4: "Obligation-based", 40: "Cash-based"}),
                     row("S", base + 3, {c: 1983 + c - 4 for c in range(4, 49)}),
                     row("S", base + 4, {47: "Program", 48: "Proposed"})])
        for i, label in enumerate(sectors + ["TOTAL"]):
            rows.append(row("S", base + 6 + i, {1: label, **{c: 6 if label == "TOTAL" else 1 for c in range(4, 49)}}))
    controls = Controls()
    records = list(sector(rows, controls, {}))
    assert len(records) == 7 * 7 * 45
    assert all(c["passed"] for c in controls.checks)
    actual = next(r for r in records if r["year"] == 2025 and r["measure"] == "expenditure")
    assert actual["status"] == "actual" and actual["reporting_basis"] == "Cash-based"
    assert actual["unit"] == "million_PHP" and actual["value_php"] == "1000000"
    assert {r["status"] for r in records if r["year"] == 2027} == {"proposed"}


def test_grdp_growth_and_indices_reconcile_and_header_repair_is_documented():
    rows = []
    for section in range(1, 7):
        base = section * 10
        rows.extend([row("S", base, {1: f"Table 2.{section} Per Capita"}),
                     row("S", base + 1, {1: "Annual 2023 to 2025"}),
                     row("S", base + 2, {1: "At Current Prices" if section % 2 else "At Constant 2018 Prices"})])
        headers = {3: 2023, 4: 2024, 5: 2025} if section not in {3, 4} else {3: "2023 - 2024", 4: "2024 - 2025"}
        if section == 6:
            headers[3] = "Annual"
        rows.append(row("S", base + 3, headers))
        for offset, entity in enumerate(["Region A", "Philippines"]):
            multiplier = 2 if entity == "Region A" else 1
            data = ({3: 100 * multiplier, 4: 110 * multiplier, 5: 121 * multiplier} if section <= 2 else
                    {3: 10, 4: 10} if section <= 4 else {3: multiplier * 100, 4: multiplier * 100, 5: multiplier * 100})
            rows.append(row("S", base + 4 + offset, {1: "A", 2: entity, **data}))
    controls = Controls()
    records = list(grdp(rows, controls, {}))
    assert len(records) == 32
    assert all(c["passed"] for c in controls.checks)
    assert any("C183" in note for note in controls.notes)


def test_population_projection_totals_and_detail_do_not_share_keys():
    rows = []
    for sheet in ["Table 1", "Table 3", "Table 4"]:
        label_column = 2 if sheet == "Table 4" else 1
        first = label_column + 1
        rows.extend([row(sheet, 1, {label_column: "Projected Population"}),
                     row(sheet, 3, {label_column: "Region, Age Group"}),
                     row(sheet, 4, {first: "01 July 2024"})])
        if sheet == "Table 4":
            rows.extend([row(sheet, 7, {1: "PH", 2: "PHILIPPINES", 3: 1000, 4: 500, 5: 500}),
                         row(sheet, 8, {1: "R", 2: "REGION A", 3: 1000, 4: 500, 5: 500})])
        else:
            rows.extend([row(sheet, 7, {1: "PHILIPPINES"}), row(sheet, 8, {1: "Total", 2: 1000, 3: 500, 4: 500}),
                         row(sheet, 9, {1: "0 - 4", 2: 100, 3: 50, 4: 50}),
                         row(sheet, 10, {1: "REGION A"}), row(sheet, 11, {1: "Total", 2: 1000, 3: 500, 4: 500})])
    controls = Controls()
    records = list(population(rows, controls, {}))
    assert len({r["record_key"] for r in records}) == len(records) == 24
    assert all(r["status"] == "projected" for r in records)
    assert all(c["passed"] for c in controls.checks)
    assert not controls.blockers


def test_agency_funding_breakdowns_are_separate_and_undefined_rate_is_not_zero():
    def amounts(label, amount):
        return {1: label, **{col: amount if (col - 2) % 5 in {0, 4} else 0 for col in range(2, 32)}}
    rows = [row("SUMMARY", 1, {1: "FY 2024 STATEMENT"}), row("SUMMARY", 3, {1: "In Thousand Pesos"}),
            row("SUMMARY", 4, {17: "ALLOTMENTS", 22: "OBLIGATIONS"}),
            row("SUMMARY", 8, {**amounts("Agency A", 10), 32: "#DIV/0!"}),
            row("SUMMARY", 46, amounts("A. TOTAL NATIONAL GOVERNMENT AGENCIES", 10)),
            row("SUMMARY", 47, amounts("Current Year Budget", 10)),
            row("SUMMARY", 50, amounts("B. SPECIAL PURPOSE FUNDS (SPFs)", 20)),
            row("SUMMARY", 51, amounts("Current Year Budget", 20))]
    controls = Controls()
    records = list(agency(rows, controls, {}))
    assert len({r["record_key"] for r in records}) == len(records)
    assert all(not r["analytical_eligible"] for r in records if r["entity"] == "Current Year Budget")
    rate = next(r for r in records if r["measure"] == "obligation_rate")
    assert rate["value_status"] == "undefined_rate" and rate["value"] == ""
    assert rate["raw_value"] == json.dumps("#DIV/0!")


def test_duplicate_source_years_are_blocked():
    results = [dict(source=f"02_approved_gaa/{name}.xlsx", years=[2024], source_sha256="same",
                    status="PASS", blockers=[]) for name in ["original", "copy"]]
    long_form.check_source_coverage(results)
    assert all(r["status"] == "REVIEW_REQUIRED" for r in results)


def test_publication_rolls_back_on_invalid_csv(gaa_file, tmp_path):
    raw, _ = gaa_file
    output = tmp_path / "out"
    result = long_form.run(raw, output, gaa_unit="PHP", gaa_unit_evidence="Synthetic fixture")
    directory = Path(result["output_directory"])
    csv_file = next(directory.glob("*-approved_budget.csv"))
    csv_file.write_text("wrong_header\nwrong_value\n")
    with pytest.raises(KeyError):
        long_form.publish_sqlite(output, directory, result)
    assert count(output) == (1, 0)
