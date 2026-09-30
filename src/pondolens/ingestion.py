"""Idempotent raw preservation and structural extraction, not analytical Silver."""

import argparse
from contextlib import closing
from datetime import datetime, timezone
import hashlib
import json
import os
from pathlib import Path
import sqlite3
import tempfile

import openpyxl
import pdfplumber

VERSION = "raw-structure-v1"
EXPECTED = {
    "01_agency_actuals": (".xlsx", 2),
    "02_approved_gaa": (".xlsx", 2),
    "03_sector_actuals": (".xlsx", 1),
    "04_regional_actuals": (".pdf", 2),
    "05_population": (".xlsx", 1),
    "06_development_indicator": (".xlsx", 1),
}


def digest(path):
    checksum = hashlib.sha256()
    with Path(path).open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            checksum.update(block)
    return checksum.hexdigest()


def json_value(value):
    # Preserve Excel types and representations rather than coercing financial data.
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def excel_records(path):
    formula_book = openpyxl.load_workbook(path, read_only=True, data_only=False)
    value_book = None
    try:
        value_book = openpyxl.load_workbook(path, read_only=True, data_only=True)
        for sheet in formula_book:
            # Reset dimensions to avoid truncation from incorrect producer metadata.
            sheet.reset_dimensions()
            cached = value_book[sheet.title]
            cached.reset_dimensions()
            yield {"kind": "sheet", "sheet": sheet.title}
            for row_number, (row, values) in enumerate(
                zip(sheet.iter_rows(), cached.iter_rows(), strict=True), start=1
            ):
                for column, (cell, value) in enumerate(zip(row, values, strict=True), start=1):
                    if cell.value is None:
                        continue
                    formula = cell.data_type == "f"
                    yield {
                        "kind": "cell", "sheet": sheet.title,
                        "row": row_number, "column": column,
                        "excel_type": cell.data_type,
                        "original_value": json_value(cell.value),
                        "cached_value": json_value(value.value) if formula else None,
                        "missing_formula_cache": formula and value.value is None,
                        "number_format": cell.number_format,
                    }
    finally:
        formula_book.close()
        if value_book is not None:
            value_book.close()


def pdf_records(path):
    with pdfplumber.open(path) as document:
        for number, page in enumerate(document.pages, start=1):
            words = page.extract_words()
            yield {
                "kind": "pdf_page", "page": number,
                "width": page.width, "height": page.height,
                "text": page.extract_text() or "",
                "words": words,
                "requires_ocr_review": not words,
            }
            for index, table in enumerate(page.find_tables(), start=1):
                yield {
                    "kind": "pdf_table_candidate", "page": number,
                    "table": index, "bbox": list(table.bbox),
                    "rows": table.extract(), "validated": False,
                }


def discover(root):
    sources = []
    for folder, (extension, count) in EXPECTED.items():
        directory = root / folder
        files = sorted(p for p in directory.glob("*") if p.is_file()
                       and p.suffix.lower() == extension and not p.name.startswith("~$"))
        if len(files) != count:
            raise ValueError(f"{folder}: expected {count} {extension} files, found {len(files)}")
        sources.extend(files)
    return sources


def ingest_one(source, raw_root, output_root, db):
    source_id = source.relative_to(raw_root).as_posix()
    archive_dir = output_root / "originals"
    archive_dir.mkdir(parents=True, exist_ok=True)

    # Snapshot first; extraction uses the archived bytes, not a changing input.
    descriptor, temporary = tempfile.mkstemp(dir=archive_dir)
    snapshot = Path(temporary)
    try:
        with os.fdopen(descriptor, "wb") as destination, source.open("rb") as original:
            for block in iter(lambda: original.read(1024 * 1024), b""):
                destination.write(block)
        file_hash = digest(snapshot)
        archive = archive_dir / f"{file_hash}{source.suffix.lower()}"
        if archive.exists():
            if digest(archive) != file_hash:
                raise ValueError(f"Archived original corrupted: {archive}")
        else:
            os.replace(snapshot, archive)
    finally:
        snapshot.unlink(missing_ok=True)

    key = hashlib.sha256(f"{source_id}|{file_hash}|{VERSION}".encode()).hexdigest()
    result = output_root / "structure" / f"{key}.jsonl"
    previous = db.execute(
        "SELECT status, output_hash FROM runs WHERE run_id = ?", (key,)
    ).fetchone()
    if previous and previous[0] == "extracted" and result.exists():
        if digest(result) == previous[1]:
            return {"source": source_id, "status": "skipped", "run_id": key}

    result.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary = tempfile.mkstemp(dir=result.parent)
    pending = Path(temporary)
    counts = {}
    try:
        records = excel_records(archive) if source.suffix.lower() == ".xlsx" else pdf_records(archive)
        with os.fdopen(descriptor, "w", encoding="utf-8") as stream:
            for record in records:
                counts[record["kind"]] = counts.get(record["kind"], 0) + 1
                record.update(source=source_id, sha256=file_hash, parser_version=VERSION)
                stream.write(json.dumps(record, ensure_ascii=False, allow_nan=False) + "\n")
        os.replace(pending, result)
        status, error, output_hash = "extracted", None, digest(result)
    except Exception as exc:
        status, error, output_hash = "failed", str(exc), None
    finally:
        pending.unlink(missing_ok=True)

    with db:
        db.execute("""
            INSERT INTO runs VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            ON CONFLICT(run_id) DO UPDATE SET
              status=excluded.status, output_hash=excluded.output_hash,
              counts=excluded.counts, error=excluded.error, updated_at=excluded.updated_at
        """, (key, source_id, file_hash, VERSION, status, output_hash,
              json.dumps(counts), error, datetime.now(timezone.utc).isoformat()))
    if error:
        raise RuntimeError(f"{source_id}: {error}")
    return {"source": source_id, "status": status, "counts": counts, "run_id": key}


def run(raw_root, output_root):
    raw_root, output_root = Path(raw_root).resolve(), Path(output_root).resolve()
    if output_root == raw_root or raw_root in output_root.parents:
        raise ValueError("Output must be outside the raw input directory")
    sources = discover(raw_root)
    output_root.mkdir(parents=True, exist_ok=True)
    summary = []
    with closing(sqlite3.connect(output_root / "audit.sqlite")) as db:
        db.execute("""CREATE TABLE IF NOT EXISTS runs (
            run_id TEXT PRIMARY KEY, source_id TEXT, file_hash TEXT, version TEXT,
            status TEXT, output_hash TEXT, counts TEXT, error TEXT, updated_at TEXT
        )""")
        for source in sources:
            try:
                summary.append(ingest_one(source, raw_root, output_root, db))
            except Exception as exc:
                summary.append({"source": source.relative_to(raw_root).as_posix(),
                                "status": "failed", "error": str(exc)})
    descriptor, temporary = tempfile.mkstemp(dir=output_root)
    with os.fdopen(descriptor, "w", encoding="utf-8") as stream:
        json.dump({"ready_for_analysis": False, "sources": summary}, stream, indent=2)
    os.replace(temporary, output_root / "latest_run.json")
    return summary


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--raw", default="data/raw")
    parser.add_argument("--output", default="outputs")
    args = parser.parse_args()
    try:
        summary = run(args.raw, args.output)
    except (ValueError, OSError) as exc:
        parser.exit(1, f"ERROR: {exc}\n")
    for item in summary:
        print(f"{item['status'].upper()}: {item['source']}")
        if item.get("error"):
            print(item["error"])
    if any(item["status"] == "failed" for item in summary):
        raise SystemExit(1)
    print("Originals and structural extraction saved. Analytical mappings are pending.")


if __name__ == "__main__":
    main()
