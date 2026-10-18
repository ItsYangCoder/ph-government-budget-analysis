"""PondoLens Google Drive -> Bronze Delta ingestion (regular Python script).

Run as a Databricks Job Python script task; set maximum concurrent runs to 1.
Configure dependencies: openpyxl, google-api-python-client, google-auth,
databricks-sdk. Catalog, schema and managed Volume are created if missing
on load runs, subject to the run identity's Unity Catalog permissions.
Default: --dataset all --dry_run true --auth_mode secret.
Existing secret: pondolens / gdrive-service-account. Optional --auth_mode file
uses --service_account_path; credentials are never stored in this code.

All years, existing Bronze table names, source strings and audit fields are
preserved. Each changed dataset is replaced from ALL its current Drive files.
Missing files or removed columns block replacement unless explicitly allowed.
Commits are atomic per Delta table, not across the whole run; retry partial runs.
Driver-side parsing is intended for small reports within configured limits.
Geometry remains raw JSON in Bronze. Regular .py file; no notebook cells.
Missing third-party dependencies are installed before initialization.
Interactive runs use native Databricks widgets when available. Explicit CLI
arguments take precedence and remain supported for Job Python script tasks.
"""

import argparse
import csv
import hashlib
import io
import json
import re
import sys
import uuid
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path

DATASET_NAMES = ['agency_actuals', 'regional_appropriation', 'regional_obligations', 'population', 'unemployment_rate', 'poverty_incidence', 'service_access', 'grdp_per_capita', 'completion_rate', 'region_boundaries', 'hospital']
DEFAULTS = {
    "catalog": "pondolens",
    "schema": "dev_bronze",
    "volume": "government_budget",
    "drive_folder_id": "1L_kvzPjQE60GjUvthLuwxbC3xKevHsoB",
    "auth_mode": "secret",
    "service_account_path": "/Workspace/Users/alumrhea@gmail.com/rock-loop-511408-h8-a51e825c4afe.json",
    "secret_scope": "pondolens",
    "secret_key": "gdrive-service-account",
    "dataset": "all",
    "dry_run": "true",
    "sheet_overrides_json": "{}",
    "allow_source_removal": "false",
    "allow_schema_removal": "false",
    "max_file_mb": "50",
    "max_total_mb": "200",
    "max_total_rows": "500000",
}


WIDGET_NAMES = (
    "catalog", "schema", "volume", "drive_folder_id", "secret_scope",
    "secret_key", "dataset", "dry_run", "sheet_overrides_json",
)



def ensure_dependencies():
    """Install missing third-party packages in the current Python environment."""
    import importlib
    import importlib.util
    import subprocess

    required = {
        "openpyxl": "openpyxl>=3.1,<4",
        "googleapiclient": "google-api-python-client>=2,<3",
        "google.auth": "google-auth>=2,<3",
        "databricks.sdk": "databricks-sdk>=0.30,<1",
    }
    missing = []
    for module, package in required.items():
        try:
            available = importlib.util.find_spec(module) is not None
        except ModuleNotFoundError:
            available = False
        if not available:
            missing.append(package)

    if not missing:
        return

    print("Installing missing dependencies: " + ", ".join(missing), flush=True)
    try:
        subprocess.run(
            [sys.executable, "-m", "pip", "install", "--disable-pip-version-check", *missing],
            check=True,
        )
    except (OSError, subprocess.CalledProcessError) as exc:
        raise RuntimeError(
            "Dependency installation failed. Configure openpyxl, "
            "google-api-python-client, google-auth and databricks-sdk "
            "in the Databricks compute/task environment, then rerun."
        ) from exc
    importlib.invalidate_caches()


def widget_parameters():
    """Use the native Databricks widget API without a separate runner notebook."""
    global dbutils
    native = globals().get("dbutils")
    if native is None:
        import builtins
        native = getattr(builtins, "dbutils", None)
    if native is None:
        try:
            from databricks.sdk.runtime import dbutils as native
        except (ImportError, RuntimeError) as exc:
            raise RuntimeError("Native dbutils is unavailable for interactive widgets. Use explicit CLI parameters for a Job run.") from exc
    widgets = getattr(native, "widgets", None)
    if widgets is None:
        raise RuntimeError("This dbutils instance has no native widgets API. No ingestion was started.")
    choices = {"dataset": ["all"] + DATASET_NAMES, "dry_run": ["true", "false"]}
    try:
        # Remove only legacy parameters created by this script, not other widgets.
        for name in DEFAULTS.keys() - set(WIDGET_NAMES):
            try:
                widgets.get(name)
            except Exception:
                continue
            widgets.remove(name)
        for name in WIDGET_NAMES:
            if name in choices:
                widgets.dropdown(name, DEFAULTS[name], choices[name])
            else:
                widgets.text(name, DEFAULTS[name])
        result = dict(DEFAULTS)
        result.update({name: widgets.get(name).strip() for name in WIDGET_NAMES})
    except (AttributeError, NotImplementedError) as exc:
        raise RuntimeError("Native widget creation is unsupported in this execution context. No ingestion was started; use the original interactive context or explicit Job CLI parameters.") from exc
    dbutils = native
    print("Parameters: Databricks widgets initialized.", flush=True)
    return result


def parameters(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    for name, default in DEFAULTS.items():
        parser.add_argument(f"--{name}", default=default)
    args = list(sys.argv[1:] if argv is None else argv)
    has_cli = any(arg.split("=", 1)[0] in {f"--{name}" for name in DEFAULTS}
                  for arg in args)
    if argv is not None or has_cli or "--help" in args or "-h" in args:
        result = vars(parser.parse_args(args))
    else:
        # Ignore unrelated interpreter arguments, as in the original script.
        result = widget_parameters()
    for name in ("dry_run", "allow_source_removal", "allow_schema_removal"):
        if result[name] not in {"true", "false"}:
            parser.error(f"{name} must be true or false")
    if result["auth_mode"] not in {"secret", "file"}:
        parser.error("auth_mode must be secret or file")
    if result["dataset"] not in ["all"] + DATASET_NAMES:
        parser.error("Unknown dataset parameter")
    if not re.fullmatch(r"[A-Za-z0-9_-]+", result["drive_folder_id"]):
        parser.error("drive_folder_id must be a folder ID, not a URL")
    for name in ("catalog", "schema", "volume"):
        try:
            identifier(result[name])
        except ValueError as exc:
            parser.error(str(exc))
    for name in ("max_file_mb", "max_total_mb", "max_total_rows"):
        try:
            result[name] = int(result[name])
        except ValueError:
            parser.error(f"{name} must be a positive integer")
        if result[name] <= 0:
            parser.error(f"{name} must be a positive integer")
    return result


def initialize(params):
    """Initialize only when executed; configure packages on the Job environment."""
    global PARAMS, CATALOG, SCHEMA, VOLUME, DRIVE_FOLDER_ID, SECRET_SCOPE, SECRET_KEY
    global DRY_RUN, SELECTED_DATASET, SHEET_OVERRIDES, RUN_ID, TARGET, ROOT
    global spark, dbutils, openpyxl, service_account, build, MediaIoBaseDownload
    global F, LongType, StringType, StructField, StructType, TimestampType
    PARAMS = params
    CATALOG, SCHEMA, VOLUME = (params[k] for k in ("catalog", "schema", "volume"))
    DRIVE_FOLDER_ID = params["drive_folder_id"]
    SECRET_SCOPE, SECRET_KEY = params["secret_scope"], params["secret_key"]
    DRY_RUN, SELECTED_DATASET = params["dry_run"] == "true", params["dataset"]
    custom = json.loads(params["sheet_overrides_json"])
    if not isinstance(custom, dict) or any(not isinstance(k, str) or not isinstance(v, str) for k, v in custom.items()):
        raise ValueError("sheet_overrides_json must map file IDs to sheet names")
    SHEET_OVERRIDES = {"1vCXocazfgtHpey2wG6sTzpQFI-pBrDfF": "Long Form", **custom}
    RUN_ID = uuid.uuid4().hex
    TARGET = f"{CATALOG}.{SCHEMA}"
    ROOT = f"/Volumes/{CATALOG}/{SCHEMA}/{VOLUME}"
    try:
        import openpyxl
        from google.oauth2 import service_account
        from googleapiclient.discovery import build
        from googleapiclient.http import MediaIoBaseDownload
        from pyspark.sql import SparkSession, functions as F
        from pyspark.sql.types import LongType, StringType, StructField, StructType, TimestampType
        from databricks.sdk import WorkspaceClient
    except ImportError as exc:
        raise RuntimeError("Configure Job dependencies: openpyxl, google-api-python-client, google-auth, databricks-sdk. Run on Databricks Spark compute.") from exc
    spark = SparkSession.builder.getOrCreate()
    if globals().get("dbutils") is None:
        dbutils = WorkspaceClient().dbutils


FOLDER_MIME = "application/vnd.google-apps.folder"
SHEET_MIME = "application/vnd.google-apps.spreadsheet"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PIPELINE_VERSION = "gdrive-snapshot-1"


def identifier(value):
    if not re.fullmatch("[A-Za-z_][A-Za-z0-9_]*", value):
        raise ValueError(f"Invalid identifier: {value!r}")
    return value



def folder_label(value):
    """Ignore numeric prefixes so folder renumbering does not break routing."""
    return re.sub("^\\d+[_ -]*", "", value.strip().lower())


def route(relative_path):
    """Route known folder labels; accept new years without code changes."""
    parts = relative_path.lower().split("/")
    root, name = (folder_label(parts[0]), parts[-1])
    direct = {
        "agency_actuals": "agency_actuals",
        "sector_expenditure": "__excluded__",
        "population": "population",
    }
    if root in direct:
        return direct[root]
    if root == "regional_expenditure":
        if "gaa" in name:
            return "regional_appropriation"
        if "obligation" in name:
            return "regional_obligations"
    if root == "health" and name.startswith(("doh_regulated_hospitals", "hospital")):
        return "hospital"
    if root == "development_indicators" and len(parts) > 2:
        return {
            "quality_education": "__excluded__",
            "unemployment_rate": "unemployment_rate",
            "poverty_incidence": "poverty_incidence",
            "service_access": "service_access",
            "grdp": "__excluded__",
            "grdp_capita": "grdp_per_capita",
            "grdp_per_capita": "grdp_per_capita",
            "student_completion_rate": "completion_rate",
            "completion_rate": "completion_rate",
        }.get(folder_label(parts[1]))
    if root in {"geographic_reference", "geographic_reference(optional)"}:
        if name.endswith((".geojson", ".json")):
            return "region_boundaries"
    return None


def source_extension(item):
    if item["mimeType"] == SHEET_MIME:
        return ".xlsx"
    return Path(item["name"].lower()).suffix


def supported_source(item, dataset):
    extension = source_extension(item)
    if dataset == "region_boundaries":
        return extension in {".geojson", ".json"}
    return extension in {".csv", ".xlsx", ".xlsm"}


def read_geojson(payload):
    """Preserve each feature and geometry as JSON strings in Bronze."""
    data = json.loads(payload.decode("utf-8-sig"))
    if not isinstance(data, dict) or data.get("type") != "FeatureCollection":
        raise ValueError("Expected a GeoJSON FeatureCollection")
    features = data.get("features")
    if not isinstance(features, list) or not features:
        raise ValueError("GeoJSON has no features")
    headers = [
        "adm1_psgc",
        "adm1_en",
        "geometry_json",
        "properties_json",
        "feature_json",
        "_rescued_data",
    ]
    records, seen = ([], set())
    for number, feature in enumerate(features, 1):
        if not isinstance(feature, dict) or feature.get("type") != "Feature":
            raise ValueError(f"Invalid GeoJSON feature {number}")
        properties = feature.get("properties")
        geometry = feature.get("geometry")
        if not isinstance(properties, dict):
            raise ValueError(f"Missing properties at feature {number}")
        if (
            not isinstance(geometry, dict)
            or geometry.get("type") not in {"Polygon", "MultiPolygon"}
            or (not geometry.get("coordinates"))
        ):
            raise ValueError(f"Missing polygon geometry at feature {number}")
        code, name = (properties.get("adm1_psgc"), properties.get("adm1_en"))
        if (
            isinstance(code, bool)
            or not re.fullmatch("\\d{1,10}", str(code))
            or (not isinstance(name, str))
            or (not name.strip())
        ):
            raise ValueError(f"Missing/invalid region code or name at feature {number}")
        key = str(code).zfill(10)
        if key in seen:
            raise ValueError(f"Duplicate region boundary: {code}")
        seen.add(key)
        encode = lambda value: json.dumps(
            value, ensure_ascii=False, allow_nan=False, separators=(",", ":")
        )
        records.append(
            (
                str(code),
                name,
                encode(geometry),
                encode(properties),
                encode(feature),
                None,
            )
        )
    return (
        headers,
        records,
        None,
        dict(
            zip(
                headers[:-1],
                [
                    "properties.adm1_psgc",
                    "properties.adm1_en",
                    "geometry",
                    "properties",
                    "feature",
                ],
            )
        ),
    )


def list_children(drive, folder_id):
    token = None
    while True:
        response = (
            drive.files()
            .list(
                q=f"'{folder_id}' in parents and trashed = false",
                pageSize=1000,
                pageToken=token,
                fields="nextPageToken,files(id,name,mimeType,modifiedTime,size)",
                supportsAllDrives=True,
                includeItemsFromAllDrives=True,
            )
            .execute(num_retries=3)
        )
        yield from response.get("files", [])
        token = response.get("nextPageToken")
        if not token:
            break


def walk(drive, folder_id, prefix="", visited=None):
    visited = set() if visited is None else visited
    if folder_id in visited:
        raise ValueError("Repeated folder in traversal")
    visited.add(folder_id)
    for item in sorted(list_children(drive, folder_id), key=lambda x: x["name"]):
        relative = prefix + item["name"]
        if item["mimeType"] == FOLDER_MIME:
            yield from walk(drive, item["id"], relative + "/", visited)
        else:
            yield {**item, "relative_path": relative}


def download(drive, item):
    if item["mimeType"] == SHEET_MIME:
        request = drive.files().export_media(fileId=item["id"], mimeType=XLSX_MIME)
    else:
        request = drive.files().get_media(fileId=item["id"], supportsAllDrives=True)
    limit = PARAMS["max_file_mb"] * 1024 * 1024
    if item.get("size") and int(item["size"]) > limit:
        raise ValueError(f"File exceeds max_file_mb: {item['relative_path']}")
    buffer = io.BytesIO()
    downloader = MediaIoBaseDownload(buffer, request, chunksize=1024 * 1024)
    done = False
    while not done:
        _, done = downloader.next_chunk(num_retries=3)
        if buffer.tell() > limit:
            raise ValueError(f"Download exceeds max_file_mb: {item['relative_path']}")
    after = (
        drive.files()
        .get(fileId=item["id"], fields="modifiedTime", supportsAllDrives=True)
        .execute(num_retries=3)
    )
    if after["modifiedTime"] != item["modifiedTime"]:
        raise RuntimeError(f"File changed during download: {item['relative_path']}")
    return buffer.getvalue()


def clean_header(row):
    result = []
    for index, value in enumerate(row):
        label = str(value).strip() if value is not None else ""
        if label.startswith("_"):
            raise ValueError(f"Reserved source header: {label}")
        label = re.sub("[^A-Za-z0-9_]", "_", label).strip("_").lower()
        label = label or f"source_column_{index + 1}"
        if label[0].isdigit():
            label = "c_" + label
        if label.startswith("_") or label in result:
            raise ValueError(f"Duplicate/reserved header: {label}")
        result.append(label)
    return result


def bounded_rows(rows):
    result = []
    limit = PARAMS["max_total_rows"] + 20
    for row in rows:
        if len(result) >= limit:
            raise ValueError("Source rows exceed max_total_rows; split the source or review the limit")
        result.append(row)
    return result


def read_rows(item, payload):
    name = item["name"].lower()
    if name.endswith((".geojson", ".json")):
        return read_geojson(payload)
    if name.endswith(".csv"):
        rows = bounded_rows(csv.reader(io.StringIO(payload.decode("utf-8-sig")), strict=True))
        sheet = None
    else:
        workbook = openpyxl.load_workbook(
            io.BytesIO(payload), read_only=True, data_only=True
        )
        try:
            sheet = SHEET_OVERRIDES.get(item["id"])
            if not sheet:
                candidates = []
                for ws in workbook.worksheets:
                    header = next(ws.iter_rows(values_only=True), ())
                    labels = {str(v).strip().lower() for v in header if v is not None}
                    long_name = re.sub("[^a-z]", "", ws.title.lower()) in {
                        "longform",
                        "longdata",
                        "populationlong",
                        "graduateslong",
                    }
                    year_header = bool(
                        labels
                        & {"year", "fiscal_year", "reference_year", "academic_year"}
                    )
                    if long_name or year_header:
                        candidates.append(ws.title)
                if len(candidates) != 1:
                    raise ValueError(
                        f"Select a tab in SHEET_OVERRIDES for {item['name']}; tabs={workbook.sheetnames}, candidates={candidates}"
                    )
                sheet = candidates[0]
            rows = bounded_rows(workbook[sheet].iter_rows(values_only=True))
        finally:
            workbook.close()
    if not rows:
        raise ValueError(f"Empty source: {item['name']}")
    header_offset = 0
    if route(item["relative_path"]) == "completion_rate":
        candidates = [
            i
            for i, row in enumerate(rows[:20])
            if {"indicator", "geolocation", "level of education", "sex"}.issubset(
                {str(v).strip().lower() for v in row if v is not None}
            )
        ]
        if len(candidates) != 1:
            raise ValueError(
                "Completion rate: expected one PSA header in first 20 rows"
            )
        header_offset = candidates[0]
        rows = rows[header_offset:]
    headers = clean_header(rows[0])
    records = []
    for number, row in enumerate(rows[1:], start=header_offset + 2):
        if not any((value is not None and str(value) != "" for value in row)):
            continue
        values = tuple((None if value is None else str(value) for value in row))
        rescued = None
        if len(values) != len(headers):
            rescued = json.dumps(
                {
                    "reason": "column_count_mismatch",
                    "source_file": item["relative_path"],
                    "source_row": number,
                    "expected_columns": len(headers),
                    "actual_columns": len(values),
                    "headers": headers,
                    "original_values": values,
                },
                ensure_ascii=False,
            )
            values = (None,) * len(headers)
        records.append(values + (rescued,))
    if not records:
        raise ValueError(f"Source has no data rows: {item['name']}")
    return (
        headers + ["_rescued_data"],
        records,
        sheet,
        dict(zip(headers, map(str, rows[0]))),
    )


def log(
    dataset,
    status,
    count,
    message=None,
    files_loaded=0,
    started_at=None,
    rescued_rows=0,
):
    record = {
        "run_id": RUN_ID,
        "dataset": dataset,
        "status": status,
        "rows_loaded": str(count),
        "finished_at": datetime.now(timezone.utc).isoformat(),
        "message": message or "",
        "files_loaded": files_loaded,
        "rescued_rows": rescued_rows,
        "started_at": started_at,
    }
    audit_types = {
        "files_loaded": LongType(),
        "rescued_rows": LongType(),
        "started_at": TimestampType(),
    }
    schema = StructType(
        [StructField(key, audit_types.get(key, StringType()), True) for key in record]
    )
    spark.createDataFrame([tuple(record.values())], schema).write.format("delta").mode(
        "append"
    ).option("mergeSchema", "true").saveAsTable(f"{TARGET}.gdrive_ingestion_log")


def load_credentials():
    """Use an existing secret or an explicitly selected credentials file."""
    if PARAMS["auth_mode"] == "secret":
        try:
            raw = dbutils.secrets.get(scope=SECRET_SCOPE, key=SECRET_KEY)
        except Exception:
            raise RuntimeError(
                "Cannot read service-account secret. Verify secret_scope, secret_key "
                "and READ access, or select auth_mode=file with service_account_path."
            ) from None
    else:
        credential_path = Path(PARAMS["service_account_path"])
        if not credential_path.is_absolute() or not credential_path.is_file():
            raise ValueError("service_account_path must point to an existing absolute JSON file path")
        raw = credential_path.read_text(encoding="utf-8-sig")
    try:
        info = json.loads(raw)
    except (ValueError, TypeError):
        raise ValueError("Credentials must contain valid service-account JSON") from None
    required = ("client_email", "private_key", "token_uri")
    if not isinstance(info, dict) or info.get("type") != "service_account" or not all(
        isinstance(info.get(key), str) and info[key].strip() for key in required
    ):
        raise ValueError("Credentials must be a service-account JSON with client_email, private_key and token_uri")
    try:
        return service_account.Credentials.from_service_account_info(
            info, scopes=["https://www.googleapis.com/auth/drive.readonly"]
        )
    except Exception:
        raise ValueError("Cannot initialize service-account credentials; check the JSON key file") from None


def validate_replacement(previous_ids, current_ids, previous_columns, new_columns):
    """Block accidental data/schema loss before snapshot replacement."""
    removed = set(previous_ids) - set(current_ids)
    if removed and PARAMS["allow_source_removal"] != "true":
        raise ValueError(f"Existing source files missing from Drive: {sorted(removed)}. Restore them or explicitly set --allow_source_removal true.")
    removed_columns = {c for c in previous_columns if not c.startswith("_")} - set(new_columns)
    if removed_columns and PARAMS["allow_schema_removal"] != "true":
        raise ValueError(f"Source columns would be removed: {sorted(removed_columns)}. Review or explicitly set --allow_schema_removal true.")


def source_snapshot(sources):
    return sorted((s["id"], s["relative_path"], s["modifiedTime"], s["mimeType"])
                  for s in sources if SELECTED_DATASET == "all" or route(s["relative_path"]) == SELECTED_DATASET)


def table_exists(table):
    """Treat a missing namespace or concurrently dropped table as absent."""
    try:
        if not spark.catalog.tableExists(table):
            return False
        # Catalog metadata can be stale right after a drop; force resolution now.
        spark.table(table).schema
        return True
    except Exception as exc:
        get_condition = getattr(exc, "getCondition", None) or getattr(exc, "getErrorClass", None)
        condition = get_condition() if get_condition else None
        if condition in {"CATALOG_NOT_FOUND", "SCHEMA_NOT_FOUND", "NO_SUCH_CATALOG_EXCEPTION", "NO_SUCH_DATABASE", "TABLE_OR_VIEW_NOT_FOUND"}:
            return False
        raise


def ensure_storage():
    """Provision Unity Catalog storage only on a load run."""
    if DRY_RUN:
        print("Dry run: catalog/schema/volume creation skipped.", flush=True)
        return
    spark.sql(f"CREATE CATALOG IF NOT EXISTS `{CATALOG}`")
    spark.sql(f"CREATE SCHEMA IF NOT EXISTS `{CATALOG}`.`{SCHEMA}`")
    spark.sql(f"CREATE VOLUME IF NOT EXISTS `{CATALOG}`.`{SCHEMA}`.`{VOLUME}`")
    dbutils.fs.ls(ROOT)


def preflight(plans, sources, drive):
    """Validate the full plan before writing landing files or Delta tables."""
    if source_snapshot(sources) != source_snapshot(list(walk(drive, DRIVE_FOLDER_ID))):
        raise RuntimeError("Drive inventory changed during validation; rerun with stable source files")
    for dataset, entries in plans.items():
        table = f"{TARGET}.{identifier(dataset)}"
        columns = {c for entry in entries for c in entry[2]}
        current_ids = {entry[0]["id"] for entry in entries}
        if dataset == "region_boundaries":
            codes = [str(row[0]).zfill(10) for entry in entries for row in entry[3]]
            if len(codes) != len(set(codes)):
                raise ValueError("Duplicate region boundaries across files; select one boundary version per dataset")
        if table_exists(table):
            previous = spark.table(table)
            if "_drive_file_id" not in previous.columns:
                raise ValueError(f"{table}: missing source lineage; review before replacing")
            ids = [r[0] for r in previous.select("_drive_file_id").distinct().collect()]
            validate_replacement(ids, current_ids, previous.columns, columns)
    if SELECTED_DATASET == "all" and PARAMS["allow_source_removal"] != "true":
        for dataset in set(DATASET_NAMES) - set(plans):
            table = f"{TARGET}.{identifier(dataset)}"
            if table_exists(table) and spark.table(table).limit(1).count():
                raise ValueError(f"No Drive files found for existing dataset {dataset}; restore sources, select one dataset, or review --allow_source_removal true")


def main():
    print(f"Starting ingestion; dataset={SELECTED_DATASET}; dry_run={DRY_RUN}", flush=True)
    credentials = load_credentials()
    drive = build("drive", "v3", credentials=credentials, cache_discovery=False)
    folder = drive.files().get(
        fileId=DRIVE_FOLDER_ID, fields="id,name", supportsAllDrives=True
    ).execute(num_retries=3)
    print(f"Connected! Source folder: {folder['name']}", flush=True)
    sources = list(walk(drive, DRIVE_FOLDER_ID))
    plans = defaultdict(list)
    errors = []
    total_bytes = total_rows = 0
    for item in sources:
        dataset = route(item["relative_path"])
        if dataset == "__excluded__":
            continue
        if SELECTED_DATASET != "all" and dataset != SELECTED_DATASET:
            continue
        supported = supported_source(item, dataset)
        if not dataset or not supported:
            errors.append(f"Unmapped/unsupported source: {item['relative_path']}")
            continue
        try:
            payload = download(drive, item)
            total_bytes += len(payload)
            if total_bytes > PARAMS["max_total_mb"] * 1024 * 1024:
                raise ValueError("Source bytes exceed max_total_mb; select one dataset per run")
            headers, records, sheet, header_map = read_rows(item, payload)
            total_rows += len(records)
            if total_rows > PARAMS["max_total_rows"]:
                raise ValueError("Source records exceed max_total_rows; select one dataset per run")
            digest = hashlib.sha256(payload).hexdigest()
            plans[dataset].append(
                (item, payload, headers, records, sheet, digest, header_map)
            )
            print(
                f"VALID {dataset}: {item['relative_path']} rows={len(records)} tab={sheet}"
            )
        except Exception as exc:
            errors.append(f"{item['relative_path']}: {exc}")
            if total_bytes > PARAMS["max_total_mb"] * 1024 * 1024 or total_rows > PARAMS["max_total_rows"]:
                break
    if not sources:
        errors.append("No sources found in Drive folder")
    if not plans and (not errors):
        errors.append(f"No files found for selected dataset: {SELECTED_DATASET}")
    if errors:
        raise ValueError("Validation failed; no tables written:\n" + "\n".join(errors))
    preflight(plans, sources, drive)
    file_count = sum((len(entries) for entries in plans.values()))
    print(f"Plan: {file_count} files -> {len(plans)} tables; dry_run={DRY_RUN}")
    if DRY_RUN:
        ensure_storage()
        return
    ensure_storage()
    failures = []
    for dataset, entries in sorted(plans.items()):
        started_at = datetime.now(timezone.utc)
        files_loaded = 0
        rescued_rows = 0
        table = f"{TARGET}.{identifier(dataset)}"
        count = sum((len(entry[3]) for entry in entries))
        fingerprint = hashlib.sha256(
            json.dumps(
                [
                    (
                        entry[0]["id"],
                        entry[0]["relative_path"],
                        (
                            hashlib.sha256(
                                json.dumps(
                                    {
                                        "headers": entry[2],
                                        "records": entry[3],
                                        "sheet": entry[4],
                                        "header_map": entry[6],
                                    },
                                    ensure_ascii=False,
                                    sort_keys=True,
                                    separators=(",", ":"),
                                ).encode("utf-8")
                            ).hexdigest()
                            if entry[0]["mimeType"] == SHEET_MIME
                            else entry[5]
                        ),
                        entry[4],
                    )
                    for entry in sorted(entries, key=lambda e: e[0]["id"])
                ]
                + [
                    (
                        PIPELINE_VERSION + "-geojson-1"
                        if dataset == "region_boundaries"
                        else PIPELINE_VERSION
                    )
                ],
                sort_keys=True,
            ).encode()
        ).hexdigest()
        try:
            if table_exists(table):
                current = spark.table(table)
                if "_dataset_hash" in current.columns:
                    hashes = [
                        r[0]
                        for r in current.select("_dataset_hash").distinct().collect()
                    ]
                    if (
                        "_rescued_data" in current.columns
                        and hashes == [fingerprint]
                        and (current.count() == count)
                    ):
                        print(f"SKIPPED {table}: unchanged source content")
                        log(
                            dataset, "SKIPPED", 0, files_loaded=0, started_at=started_at
                        )
                        continue
            combined = None
            for item, payload, headers, records, sheet, digest, header_map in entries:
                extension = source_extension(item)
                path = Path(ROOT) / "gdrive_landing" / item["id"] / (digest + extension)
                path.parent.mkdir(parents=True, exist_ok=True)
                path.write_bytes(payload)
                schema = StructType(
                    [StructField(c, StringType(), True) for c in headers]
                )
                frame = spark.createDataFrame(records, schema)
                for column, value in {
                    "_drive_file_id": item["id"],
                    "_source_file": item["relative_path"],
                    "_landing_file": str(path),
                    "_source_modified_at": item["modifiedTime"],
                    "_source_sheet": sheet,
                    "_content_hash": digest,
                    "_header_map": json.dumps(header_map),
                    "_dataset_hash": fingerprint,
                    "_run_id": RUN_ID,
                }.items():
                    frame = frame.withColumn(column, F.lit(value).cast("string"))
                frame = frame.withColumn("_ingested_at", F.current_timestamp())
                combined = (
                    frame
                    if combined is None
                    else combined.unionByName(frame, allowMissingColumns=True)
                )
            if combined.count() != count:
                raise RuntimeError(f"Pre-write row count mismatch for {table}")
            combined.write.format("delta").mode("overwrite").option(
                "overwriteSchema", "true"
            ).saveAsTable(table)
            files_loaded = len(entries)
            actual = spark.table(table).count()
            rescued_rows = (
                spark.table(table).where(F.col("_rescued_data").isNotNull()).count()
            )
            if actual != count:
                raise RuntimeError(
                    f"Row count mismatch: expected={count}, actual={actual}"
                )
            log(
                dataset,
                "SUCCESS",
                actual,
                files_loaded=files_loaded,
                started_at=started_at,
                rescued_rows=rescued_rows,
            )
            print(f"SUCCESS {table}: files={len(entries)} rows={actual}")
        except Exception as exc:
            failures.append(f"{dataset}: {exc}")
            try:
                log(
                    dataset,
                    "FAILED",
                    0,
                    str(exc)[:2000],
                    files_loaded=files_loaded,
                    started_at=started_at,
                    rescued_rows=rescued_rows,
                )
            except Exception as log_error:
                print(f"Log failed for {dataset}: {log_error}")
    if failures:
        raise RuntimeError(
            "Some datasets failed; rerun after fixing:\n" + "\n".join(failures)
        )
    print(f"Completed run {RUN_ID}. Log: {TARGET}.gdrive_ingestion_log")


if __name__ == "__main__":
    ensure_dependencies()
    initialize(parameters())
    main()
